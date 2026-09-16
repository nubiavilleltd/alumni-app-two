"""Bounded parsing for administrator-supplied alumni roster imports."""

from __future__ import annotations

import csv
import json
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass
from datetime import UTC, date, datetime
from io import BytesIO, StringIO
from pathlib import Path, PurePosixPath
from typing import Any
from zipfile import BadZipFile, LargeZipFile, ZipFile

from email_validator import EmailNotValidError, validate_email
from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.utils.exceptions import InvalidFileException  # type: ignore[import-untyped]

MAX_ALUMNI_IMPORT_BYTES = 2 * 1024 * 1024
MAX_ALUMNI_IMPORT_ROWS = 500
MAX_ALUMNI_ARCHIVE_BYTES = 20 * 1024 * 1024
MAX_ALUMNI_ARCHIVE_FILES = 100
MAX_ALUMNI_ARCHIVE_RATIO = 200
MAX_ALUMNI_MULTIPART_BYTES = MAX_ALUMNI_IMPORT_BYTES + 128 * 1024

_CSV_CONTENT_TYPES = frozenset(
    {
        "",
        "application/csv",
        "application/octet-stream",
        "application/vnd.ms-excel",
        "text/csv",
        "text/plain",
    }
)
_XLSX_CONTENT_TYPES = frozenset(
    {
        "",
        "application/octet-stream",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/zip",
    }
)
_HEADER_ALIASES = {
    "email": "email",
    "email address": "email",
    "surname": "last_name",
    "last name": "last_name",
    "other names": "other_names",
    "first name": "other_names",
    "name in school (first name + surname)": "name_in_school",
    "name in school": "name_in_school",
    "whatsapp mobile phone number (e.g. 080xxxxxxxx)": "phone",
    "phone": "phone",
    "alternative phone number": "alternative_phone",
    "alternative phone": "alternative_phone",
    "birth date": "birth_date",
    "date of birth": "birth_date",
    "year of graduation from fggc owerri (e.g. 1994)": "graduation_year",
    "graduation year": "graduation_year",
    "year of graduation": "graduation_year",
    "house color": "house_color",
    "house colour": "house_color",
    "are you the coordinator for your class?": "is_coordinator",
    "is coordinator": "is_coordinator",
    "residential address (please include the closest bus stop/ landmark)": ("residential_address"),
    "residential address": "residential_address",
    "area": "area",
    "city": "city",
    "current employment status": "employment_status",
    "employment status": "employment_status",
    "occupation(s)/ profession(s). fill as appropriate.": "occupation",
    "occupation": "occupation",
    "industry sector (select as many as are applicable)": "industry_sector",
    "industry sector": "industry_sector",
    "years of professional experience": "years_of_experience",
    "years of experience": "years_of_experience",
    "would you be interested in volunteering for any project/ initiative?": ("is_volunteer"),
    "is volunteer": "is_volunteer",
    "timestamp": "source_timestamp",
    "created at": "source_timestamp",
}
_REQUIRED_FIELDS = frozenset({"email", "last_name", "other_names", "graduation_year", "city"})


class AlumniImportFileError(Exception):
    """An alumni roster failed a bounded request, file, header, or row check."""

    def __init__(self, code: str, message: str, http_status: int) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.http_status = http_status


@dataclass(frozen=True, slots=True)
class AlumniImportRow:
    """One validated roster row ready for transactional reconciliation."""

    source_row: int
    email: str
    first_name: str
    last_name: str
    fullname: str
    name_in_school: str
    phone: str
    alternative_phone: str | None
    birth_date: date | None
    graduation_year: int
    house_color: str
    coordinator_requested: bool
    residential_address: str
    area: str
    city: str
    employment_status: str
    occupation: str
    industry_sector: str
    years_of_experience: str | None
    is_volunteer: bool


def parse_alumni_import_file(
    filename: str | None,
    content: bytes,
    declared_content_type: str | None = None,
) -> tuple[AlumniImportRow, ...]:
    """Parse one small CSV or XLSX roster without evaluating workbook formulas."""
    if not filename or not Path(filename).name:
        raise AlumniImportFileError(
            "alumni_import_file_required",
            "A CSV or XLSX file is required",
            400,
        )
    _require_bounded_content(content)
    extension = Path(filename).suffix.casefold()
    content_type = (declared_content_type or "").split(";", 1)[0].strip().casefold()
    if extension == ".xls":
        raise AlumniImportFileError(
            "alumni_import_legacy_xls_unsupported",
            "Legacy XLS files are not accepted; convert the file to XLSX or CSV",
            415,
        )
    if extension == ".csv":
        if content_type not in _CSV_CONTENT_TYPES:
            raise _unsupported_type()
        records = _csv_records(content)
    elif extension == ".xlsx":
        if content_type not in _XLSX_CONTENT_TYPES:
            raise _unsupported_type()
        records = _xlsx_records(content)
    else:
        raise _unsupported_type()
    return _validated_tabular_rows(records)


def parse_alumni_import_json(content: bytes) -> tuple[int, tuple[AlumniImportRow, ...]]:
    """Parse one bounded JSON request with an explicit chapter and record array."""
    _require_bounded_content(content)
    try:
        payload = json.loads(
            content.decode("utf-8"),
            parse_constant=lambda value: (_ for _ in ()).throw(ValueError(value)),
        )
    except (UnicodeDecodeError, json.JSONDecodeError, ValueError) as exc:
        raise AlumniImportFileError(
            "alumni_import_invalid_request",
            "The JSON import request is invalid",
            400,
        ) from exc
    if not isinstance(payload, dict) or set(payload) != {"chapter_id", "records"}:
        raise AlumniImportFileError(
            "alumni_import_invalid_request",
            "JSON imports require only chapter_id and records",
            400,
        )
    chapter_id = _positive_identifier(payload["chapter_id"])
    records = payload["records"]
    if not isinstance(records, list) or not records:
        raise AlumniImportFileError(
            "alumni_import_no_rows",
            "The import must contain at least one record",
            400,
        )
    if len(records) > MAX_ALUMNI_IMPORT_ROWS:
        raise AlumniImportFileError(
            "alumni_import_too_many_rows",
            f"The import must contain at most {MAX_ALUMNI_IMPORT_ROWS} records",
            413,
        )
    rows: list[AlumniImportRow] = []
    emails: dict[str, int] = {}
    for row_number, record in enumerate(records, start=1):
        if not isinstance(record, dict):
            raise _invalid_row(row_number, "must be an object")
        canonical = _canonical_mapping(record, row_number)
        row = _validated_row(row_number, canonical)
        _remember_email(emails, row)
        rows.append(row)
    return chapter_id, tuple(rows)


def _require_bounded_content(content: bytes) -> None:
    if not content:
        raise AlumniImportFileError(
            "alumni_import_empty_request",
            "The import request is empty",
            400,
        )
    if len(content) > MAX_ALUMNI_IMPORT_BYTES:
        raise AlumniImportFileError(
            "alumni_import_too_large",
            "The import request must not exceed 2 MB",
            413,
        )


def _unsupported_type() -> AlumniImportFileError:
    return AlumniImportFileError(
        "alumni_import_unsupported_type",
        "Only JSON, CSV, and XLSX alumni imports are accepted",
        415,
    )


def _csv_records(content: bytes) -> list[tuple[int, list[object]]]:
    if b"\x00" in content:
        raise AlumniImportFileError(
            "alumni_import_invalid_file",
            "The uploaded CSV file is invalid",
            400,
        )
    try:
        text = content.decode("utf-8-sig")
        first_line = text.splitlines()[0] if text.splitlines() else ""
        delimiter = ";" if first_line.count(";") > first_line.count(",") else ","
        reader = csv.reader(StringIO(text, newline=""), delimiter=delimiter, strict=True)
        return [(reader.line_num, list(row)) for row in reader]
    except (UnicodeDecodeError, csv.Error) as exc:
        raise AlumniImportFileError(
            "alumni_import_invalid_file",
            "The uploaded CSV file is invalid",
            400,
        ) from exc


def _xlsx_records(content: bytes) -> list[tuple[int, list[object]]]:
    _validate_xlsx_archive(content)
    workbook: Any | None = None
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
        sheet = workbook.active
        if sheet.max_column > len(_HEADER_ALIASES):
            raise AlumniImportFileError(
                "alumni_import_invalid_headers",
                "The spreadsheet contains too many columns",
                400,
            )
        if sheet.max_row > MAX_ALUMNI_IMPORT_ROWS + 1:
            raise AlumniImportFileError(
                "alumni_import_too_many_rows",
                f"The import must contain at most {MAX_ALUMNI_IMPORT_ROWS} records",
                413,
            )
        records: list[tuple[int, list[object]]] = []
        for row_number, cells in enumerate(sheet.iter_rows(), start=1):
            if any(cell.data_type == "f" for cell in cells):
                raise AlumniImportFileError(
                    "alumni_import_formula_not_allowed",
                    f"Spreadsheet formulas are not accepted (row {row_number})",
                    400,
                )
            records.append((row_number, [cell.value for cell in cells]))
        return records
    except AlumniImportFileError:
        raise
    except (
        BadZipFile,
        EOFError,
        InvalidFileException,
        KeyError,
        OSError,
        TypeError,
        ValueError,
    ) as exc:
        raise AlumniImportFileError(
            "alumni_import_invalid_file",
            "The uploaded XLSX file is invalid",
            400,
        ) from exc
    finally:
        if workbook is not None:
            workbook.close()


def _validate_xlsx_archive(content: bytes) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if not entries or len(entries) > MAX_ALUMNI_ARCHIVE_FILES:
                raise AlumniImportFileError(
                    "alumni_import_invalid_file",
                    "The uploaded XLSX file is invalid",
                    400,
                )
            names = {entry.filename for entry in entries}
            if not {"[Content_Types].xml", "xl/workbook.xml"}.issubset(names):
                raise AlumniImportFileError(
                    "alumni_import_invalid_file",
                    "The uploaded XLSX file is invalid",
                    400,
                )
            total_size = 0
            for entry in entries:
                parts = PurePosixPath(entry.filename).parts
                if entry.filename.startswith("/") or ".." in parts or entry.flag_bits & 0x1:
                    raise AlumniImportFileError(
                        "alumni_import_invalid_file",
                        "The uploaded XLSX file is invalid",
                        400,
                    )
                total_size += entry.file_size
                if total_size > MAX_ALUMNI_ARCHIVE_BYTES:
                    raise AlumniImportFileError(
                        "alumni_import_archive_too_large",
                        "The XLSX file expands beyond the permitted size",
                        413,
                    )
                if (
                    entry.file_size > 10_000
                    and entry.file_size > max(entry.compress_size, 1) * MAX_ALUMNI_ARCHIVE_RATIO
                ):
                    raise AlumniImportFileError(
                        "alumni_import_archive_too_large",
                        "The XLSX file expands beyond the permitted size",
                        413,
                    )
    except AlumniImportFileError:
        raise
    except (BadZipFile, LargeZipFile, OSError, ValueError) as exc:
        raise AlumniImportFileError(
            "alumni_import_invalid_file",
            "The uploaded XLSX file is invalid",
            400,
        ) from exc


def _validated_tabular_rows(
    records: Iterable[tuple[int, list[object]]],
) -> tuple[AlumniImportRow, ...]:
    headers: list[str] | None = None
    rows: list[AlumniImportRow] = []
    emails: dict[str, int] = {}
    for row_number, values in records:
        if _blank_row(values):
            continue
        if headers is None:
            headers = _canonical_headers(values)
            continue
        if len(values) != len(headers):
            raise _invalid_row(row_number, "does not match the header column count")
        row = _validated_row(row_number, dict(zip(headers, values, strict=True)))
        _remember_email(emails, row)
        rows.append(row)
        if len(rows) > MAX_ALUMNI_IMPORT_ROWS:
            raise AlumniImportFileError(
                "alumni_import_too_many_rows",
                f"The import must contain at most {MAX_ALUMNI_IMPORT_ROWS} records",
                413,
            )
    if headers is None:
        raise AlumniImportFileError(
            "alumni_import_invalid_headers",
            "The import header row is missing",
            400,
        )
    if not rows:
        raise AlumniImportFileError(
            "alumni_import_no_rows",
            "The import must contain at least one record",
            400,
        )
    return tuple(rows)


def _canonical_headers(values: Sequence[object]) -> list[str]:
    canonical = [_canonical_header(value) for value in values]
    if any(not header for header in canonical) or len(set(canonical)) != len(canonical):
        raise AlumniImportFileError(
            "alumni_import_invalid_headers",
            "The import contains unknown or duplicate headers",
            400,
        )
    missing = sorted(_REQUIRED_FIELDS.difference(canonical))
    if missing:
        raise AlumniImportFileError(
            "alumni_import_invalid_headers",
            f"The import is missing required columns: {', '.join(missing)}",
            400,
        )
    return canonical


def _canonical_mapping(record: Mapping[object, object], row_number: int) -> dict[str, object]:
    canonical: dict[str, object] = {}
    for raw_key, value in record.items():
        header = _canonical_header(raw_key)
        if not header or header in canonical:
            raise _invalid_row(row_number, "contains unknown or duplicate fields")
        canonical[header] = value
    missing = sorted(_REQUIRED_FIELDS.difference(canonical))
    if missing:
        raise _invalid_row(row_number, f"is missing required fields: {', '.join(missing)}")
    return canonical


def _canonical_header(value: object) -> str:
    if not isinstance(value, str):
        return ""
    normalized = " ".join(value.replace("_", " ").split()).casefold()
    if normalized.startswith("city "):
        return "city"
    return _HEADER_ALIASES.get(normalized, "")


def _validated_row(row_number: int, values: Mapping[str, object]) -> AlumniImportRow:
    email = _email(values.get("email"), row_number)
    other_names = _text(values.get("other_names"), "other_names", 80, row_number, required=True)
    first_name = other_names.split()[0]
    if len(first_name) > 50:
        raise _invalid_row(row_number, "has a first name longer than 50 characters")
    last_name = _text(values.get("last_name"), "last_name", 50, row_number, required=True)
    fullname = f"{other_names} {last_name}".strip()
    if len(fullname) > 100:
        raise _invalid_row(row_number, "has a full name longer than 100 characters")
    graduation_year = _graduation_year(values.get("graduation_year"), row_number)
    city = _text(values.get("city"), "city", 100, row_number, required=True)
    coordinator_requested = _boolean(values.get("is_coordinator"), "is_coordinator", row_number)
    return AlumniImportRow(
        source_row=row_number,
        email=email,
        first_name=first_name,
        last_name=last_name,
        fullname=fullname,
        name_in_school=_text(values.get("name_in_school"), "name_in_school", 200, row_number)
        or fullname,
        phone=_text(values.get("phone"), "phone", 20, row_number),
        alternative_phone=_optional_text(
            values.get("alternative_phone"), "alternative_phone", 20, row_number
        ),
        birth_date=_birth_date(values.get("birth_date"), row_number),
        graduation_year=graduation_year,
        house_color=_text(values.get("house_color"), "house_color", 50, row_number),
        coordinator_requested=coordinator_requested,
        residential_address=_text(
            values.get("residential_address"), "residential_address", 5_000, row_number
        ),
        area=_text(values.get("area"), "area", 100, row_number),
        city=city,
        employment_status=_text(
            values.get("employment_status"), "employment_status", 100, row_number
        ),
        occupation=_text(values.get("occupation"), "occupation", 5_000, row_number),
        industry_sector=_text(values.get("industry_sector"), "industry_sector", 5_000, row_number),
        years_of_experience=_optional_text(
            values.get("years_of_experience"), "years_of_experience", 50, row_number
        ),
        is_volunteer=_boolean(values.get("is_volunteer"), "is_volunteer", row_number),
    )


def _email(value: object, row_number: int) -> str:
    raw = _text(value, "email", 150, row_number, required=True).casefold()
    try:
        normalized = validate_email(raw, check_deliverability=False).normalized.casefold()
    except EmailNotValidError as exc:
        raise _invalid_row(row_number, "has an invalid email") from exc
    if len(normalized) > 150:
        raise _invalid_row(row_number, "has an email longer than 150 characters")
    return normalized


def _text(
    value: object,
    field: str,
    max_length: int,
    row_number: int,
    *,
    required: bool = False,
) -> str:
    if value is None:
        normalized = ""
    elif isinstance(value, (str, int, float)) and not isinstance(value, bool):
        normalized = " ".join(str(value).split())
    else:
        raise _invalid_row(row_number, f"has an invalid {field}")
    if required and not normalized:
        raise _invalid_row(row_number, f"is missing {field}")
    if len(normalized) > max_length or any(
        ord(character) < 32 or ord(character) == 127 for character in normalized
    ):
        raise _invalid_row(row_number, f"has an invalid {field}")
    return normalized


def _optional_text(
    value: object,
    field: str,
    max_length: int,
    row_number: int,
) -> str | None:
    return _text(value, field, max_length, row_number) or None


def _birth_date(value: object, row_number: int) -> date | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    parsed_date: date
    if isinstance(value, datetime):
        parsed_date = value.date()
    elif isinstance(value, date):
        parsed_date = value
    elif isinstance(value, str):
        candidate: date | None = None
        for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                candidate = datetime.strptime(value.strip(), pattern).replace(tzinfo=UTC).date()
                break
            except ValueError:
                continue
        if candidate is None:
            raise _invalid_row(row_number, "has an invalid birth_date")
        parsed_date = candidate
    else:
        raise _invalid_row(row_number, "has an invalid birth_date")
    if parsed_date > datetime.now(UTC).date():
        raise _invalid_row(row_number, "has a future birth_date")
    return parsed_date


def _graduation_year(value: object, row_number: int) -> int:
    if isinstance(value, bool) or not isinstance(value, (str, int, float)):
        raise _invalid_row(row_number, "has an invalid graduation_year")
    try:
        year = int(value)
    except (TypeError, ValueError, OverflowError) as exc:
        raise _invalid_row(row_number, "has an invalid graduation_year") from exc
    if isinstance(value, float) and not value.is_integer():
        raise _invalid_row(row_number, "has an invalid graduation_year")
    if year < 1966 or year > datetime.now(UTC).year:
        raise _invalid_row(row_number, "has an invalid graduation_year")
    return year


def _boolean(value: object, field: str, row_number: int) -> bool:
    if value is None or value == "":
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, int) and value in {0, 1}:
        return bool(value)
    if isinstance(value, str):
        normalized = value.strip().casefold()
        if normalized in {"yes", "true", "1", "on"}:
            return True
        if normalized in {"no", "false", "0", "off"}:
            return False
    raise _invalid_row(row_number, f"has an invalid {field}")


def _positive_identifier(value: object) -> int:
    if isinstance(value, bool) or not isinstance(value, (str, int, float)):
        raise AlumniImportFileError(
            "alumni_import_invalid_chapter",
            "chapter_id must be a positive integer",
            400,
        )
    try:
        identifier = int(value)
    except (TypeError, ValueError, OverflowError) as exc:
        raise AlumniImportFileError(
            "alumni_import_invalid_chapter",
            "chapter_id must be a positive integer",
            400,
        ) from exc
    if identifier <= 0 or str(value).strip() != str(identifier):
        raise AlumniImportFileError(
            "alumni_import_invalid_chapter",
            "chapter_id must be a positive integer",
            400,
        )
    return identifier


def parse_chapter_id(value: object) -> int:
    """Validate the multipart chapter field with the JSON contract's exact rules."""
    return _positive_identifier(value)


def _remember_email(emails: dict[str, int], row: AlumniImportRow) -> None:
    previous = emails.get(row.email)
    if previous is not None:
        raise AlumniImportFileError(
            "alumni_import_duplicate_email",
            f"Email is repeated in rows {previous} and {row.source_row}",
            400,
        )
    emails[row.email] = row.source_row


def _blank_row(values: Sequence[object]) -> bool:
    return all(value is None or (isinstance(value, str) and not value.strip()) for value in values)


def _invalid_row(row_number: int, reason: str) -> AlumniImportFileError:
    return AlumniImportFileError(
        "alumni_import_invalid_row",
        f"Row {row_number} {reason}",
        400,
    )
