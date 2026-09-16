"""Bounded parsing for administrator-supplied zone and city catalogues."""

from __future__ import annotations

import csv
from collections.abc import Iterable
from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import Path, PurePosixPath
from typing import Any
from zipfile import BadZipFile, LargeZipFile, ZipFile

from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.utils.exceptions import InvalidFileException  # type: ignore[import-untyped]

MAX_GEOGRAPHY_IMPORT_BYTES = 2 * 1024 * 1024
MAX_GEOGRAPHY_IMPORT_ROWS = 2_000
MAX_GEOGRAPHY_ARCHIVE_BYTES = 20 * 1024 * 1024
MAX_GEOGRAPHY_ARCHIVE_FILES = 100
MAX_GEOGRAPHY_ARCHIVE_RATIO = 200
MAX_GEOGRAPHY_MULTIPART_BYTES = MAX_GEOGRAPHY_IMPORT_BYTES + 128 * 1024

_CSV_CONTENT_TYPES = frozenset(
    {
        "",
        "application/csv",
        "application/octet-stream",
        "application/vnd.ms-excel",
        "text/csv",
        "text/plain",
    }
)
_XLSX_CONTENT_TYPES = frozenset(
    {
        "",
        "application/octet-stream",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/zip",
    }
)
_EXPECTED_HEADERS = frozenset({"zone", "city"})


class GeographyImportFileError(Exception):
    """An uploaded geography catalogue failed a bounded file or row check."""

    def __init__(self, code: str, message: str, http_status: int) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.http_status = http_status


@dataclass(frozen=True, slots=True)
class GeographyImportRow:
    """One validated source row ready for transactional catalogue reconciliation."""

    source_row: int
    zone: str
    city: str


def parse_geography_import(
    filename: str | None,
    content: bytes,
    declared_content_type: str | None = None,
) -> tuple[GeographyImportRow, ...]:
    """Parse one small CSV or XLSX catalogue without evaluating workbook formulas."""
    if not filename or not Path(filename).name:
        raise GeographyImportFileError(
            "geography_import_file_required",
            "A CSV or XLSX file is required",
            400,
        )
    if not content:
        raise GeographyImportFileError(
            "geography_import_empty_file",
            "The uploaded file is empty",
            400,
        )
    if len(content) > MAX_GEOGRAPHY_IMPORT_BYTES:
        raise GeographyImportFileError(
            "geography_import_too_large",
            "The uploaded file must not exceed 2 MB",
            413,
        )

    extension = Path(filename).suffix.casefold()
    content_type = (declared_content_type or "").split(";", 1)[0].strip().casefold()
    if extension == ".xls":
        raise GeographyImportFileError(
            "geography_import_legacy_xls_unsupported",
            "Legacy XLS files are not accepted; convert the file to XLSX or CSV",
            415,
        )
    if extension == ".csv":
        if content_type not in _CSV_CONTENT_TYPES:
            raise _unsupported_type()
        records = _csv_records(content)
    elif extension == ".xlsx":
        if content_type not in _XLSX_CONTENT_TYPES:
            raise _unsupported_type()
        records = _xlsx_records(content)
    else:
        raise _unsupported_type()
    return _validated_rows(records)


def _unsupported_type() -> GeographyImportFileError:
    return GeographyImportFileError(
        "geography_import_unsupported_type",
        "Only CSV and XLSX geography files are accepted",
        415,
    )


def _csv_records(content: bytes) -> list[tuple[int, list[object]]]:
    if b"\x00" in content:
        raise GeographyImportFileError(
            "geography_import_invalid_file",
            "The uploaded CSV file is invalid",
            400,
        )
    try:
        text = content.decode("utf-8-sig")
        reader = csv.reader(StringIO(text, newline=""), strict=True)
        return [(reader.line_num, list(row)) for row in reader]
    except (UnicodeDecodeError, csv.Error) as exc:
        raise GeographyImportFileError(
            "geography_import_invalid_file",
            "The uploaded CSV file is invalid",
            400,
        ) from exc


def _xlsx_records(content: bytes) -> list[tuple[int, list[object]]]:
    _validate_xlsx_archive(content)
    workbook: Any | None = None
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
        sheet = workbook.active
        if sheet.max_column > 2:
            raise GeographyImportFileError(
                "geography_import_invalid_headers",
                "The file must contain only zone and city columns",
                400,
            )
        if sheet.max_row > MAX_GEOGRAPHY_IMPORT_ROWS + 1:
            raise GeographyImportFileError(
                "geography_import_too_many_rows",
                f"The file must contain at most {MAX_GEOGRAPHY_IMPORT_ROWS} data rows",
                413,
            )
        records: list[tuple[int, list[object]]] = []
        for row_number, cells in enumerate(sheet.iter_rows(max_col=2), start=1):
            if any(cell.data_type == "f" for cell in cells):
                raise GeographyImportFileError(
                    "geography_import_formula_not_allowed",
                    f"Spreadsheet formulas are not accepted (row {row_number})",
                    400,
                )
            records.append((row_number, [cell.value for cell in cells]))
        return records
    except GeographyImportFileError:
        raise
    except (
        BadZipFile,
        EOFError,
        InvalidFileException,
        KeyError,
        OSError,
        TypeError,
        ValueError,
    ) as exc:
        raise GeographyImportFileError(
            "geography_import_invalid_file",
            "The uploaded XLSX file is invalid",
            400,
        ) from exc
    finally:
        if workbook is not None:
            workbook.close()


def _validate_xlsx_archive(content: bytes) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if not entries or len(entries) > MAX_GEOGRAPHY_ARCHIVE_FILES:
                raise GeographyImportFileError(
                    "geography_import_invalid_file",
                    "The uploaded XLSX file is invalid",
                    400,
                )
            names = {entry.filename for entry in entries}
            if not {"[Content_Types].xml", "xl/workbook.xml"}.issubset(names):
                raise GeographyImportFileError(
                    "geography_import_invalid_file",
                    "The uploaded XLSX file is invalid",
                    400,
                )
            total_size = 0
            for entry in entries:
                parts = PurePosixPath(entry.filename).parts
                if entry.filename.startswith("/") or ".." in parts or entry.flag_bits & 0x1:
                    raise GeographyImportFileError(
                        "geography_import_invalid_file",
                        "The uploaded XLSX file is invalid",
                        400,
                    )
                total_size += entry.file_size
                if total_size > MAX_GEOGRAPHY_ARCHIVE_BYTES:
                    raise GeographyImportFileError(
                        "geography_import_archive_too_large",
                        "The XLSX file expands beyond the permitted size",
                        413,
                    )
                if (
                    entry.file_size > 10_000
                    and entry.file_size > max(entry.compress_size, 1) * MAX_GEOGRAPHY_ARCHIVE_RATIO
                ):
                    raise GeographyImportFileError(
                        "geography_import_archive_too_large",
                        "The XLSX file expands beyond the permitted size",
                        413,
                    )
    except GeographyImportFileError:
        raise
    except (BadZipFile, LargeZipFile, OSError, ValueError) as exc:
        raise GeographyImportFileError(
            "geography_import_invalid_file",
            "The uploaded XLSX file is invalid",
            400,
        ) from exc


def _validated_rows(
    records: Iterable[tuple[int, list[object]]],
) -> tuple[GeographyImportRow, ...]:
    header_map: dict[str, int] | None = None
    rows: list[GeographyImportRow] = []
    city_rows: dict[str, int] = {}

    for row_number, values in records:
        if all(value is None or (isinstance(value, str) and not value.strip()) for value in values):
            continue
        if header_map is None:
            headers = [_header_name(value) for value in values]
            if len(headers) != 2 or len(set(headers)) != 2 or set(headers) != _EXPECTED_HEADERS:
                raise GeographyImportFileError(
                    "geography_import_invalid_headers",
                    "The file must contain only zone and city columns",
                    400,
                )
            header_map = {header: index for index, header in enumerate(headers)}
            continue
        if len(values) != 2:
            raise GeographyImportFileError(
                "geography_import_invalid_row",
                f"Row {row_number} must contain exactly one zone and one city",
                400,
            )
        zone = _catalogue_name(values[header_map["zone"]], "zone", 100, row_number)
        city = _catalogue_name(values[header_map["city"]], "city", 150, row_number)
        city_key = city.casefold()
        if city_key in city_rows:
            raise GeographyImportFileError(
                "geography_import_duplicate_city",
                f"City is repeated in rows {city_rows[city_key]} and {row_number}",
                400,
            )
        city_rows[city_key] = row_number
        rows.append(GeographyImportRow(source_row=row_number, zone=zone, city=city))
        if len(rows) > MAX_GEOGRAPHY_IMPORT_ROWS:
            raise GeographyImportFileError(
                "geography_import_too_many_rows",
                f"The file must contain at most {MAX_GEOGRAPHY_IMPORT_ROWS} data rows",
                413,
            )

    if header_map is None:
        raise GeographyImportFileError(
            "geography_import_invalid_headers",
            "The file must contain only zone and city columns",
            400,
        )
    if not rows:
        raise GeographyImportFileError(
            "geography_import_no_rows",
            "The file contains no geography rows",
            400,
        )
    return tuple(rows)


def _header_name(value: object) -> str:
    if not isinstance(value, str):
        return ""
    return " ".join(value.split()).casefold()


def _catalogue_name(value: object, field: str, max_length: int, row_number: int) -> str:
    if not isinstance(value, str):
        raise GeographyImportFileError(
            "geography_import_invalid_row",
            f"Row {row_number} has an invalid {field} value",
            400,
        )
    normalized = " ".join(value.split())
    if (
        not normalized
        or len(normalized) > max_length
        or any(ord(character) < 32 or ord(character) == 127 for character in normalized)
    ):
        raise GeographyImportFileError(
            "geography_import_invalid_row",
            f"Row {row_number} has an invalid {field} value",
            400,
        )
    return normalized
